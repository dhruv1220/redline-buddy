"""Draft a negotiation letter to the counterparty from review findings.

Deterministic and local-only: each finding's contract excerpt, plain-language
concern, requested change, and quotable fallback language becomes one numbered
change request. The output is a starting draft — not legal advice, not a
lawyer's judgment — meant for attorney review before anything is sent.
"""

from __future__ import annotations

import io
from datetime import date
from typing import TYPE_CHECKING

from docx import Document
from docx.shared import Inches

from .review import SEVERITY_RANK, Finding

if TYPE_CHECKING:
    from .compare import Comparison

DISCLAIMER = (
    "Not legal advice. This is a machine-generated draft based on heuristic "
    "checks, not a lawyer's judgment. Have an attorney review it before you "
    "send anything."
)

_SEVERITY_LABEL = {
    "critical": "CRITICAL",
    "high": "HIGH",
    "medium": "MEDIUM",
    "low": "LOW",
}

_RECIPIENT_PLACEHOLDER = "[Counterparty name]"
_SENDER_PLACEHOLDER = "[Your name]"


def select_findings(
    findings: list[Finding], min_severity: str = "medium"
) -> list[Finding]:
    """Keep findings at or above ``min_severity``, ordered worst-first.

    ``review_contract`` already sorts this way, but letters built from
    hand-assembled finding lists should get the same ordering.
    """
    threshold = SEVERITY_RANK.get(min_severity, 2)
    kept = [f for f in findings if SEVERITY_RANK.get(f.severity, 9) <= threshold]
    kept.sort(key=lambda f: (SEVERITY_RANK.get(f.severity, 9), f.rule_id))
    return kept


def _h1(text: str, plain: bool) -> str:
    return text if plain else f"# {text}"


def _h2(text: str, plain: bool) -> str:
    return text if plain else f"## {text}"


def _bold(text: str, plain: bool) -> str:
    return text if plain else f"**{text}**"


def _quote(text: str, plain: bool) -> str:
    lines = text.strip().splitlines() or [""]
    if plain:
        return "\n".join(f"    {line}" for line in lines)
    return "\n".join(f"> {line}" if line.strip() else ">" for line in lines)


def _disclaimer(plain: bool) -> str:
    return DISCLAIMER if plain else f"> **Not legal advice.** {DISCLAIMER}"


def _render_no_findings(
    contract_name: str, playbook_name: str, plain: bool
) -> str:
    lines = [
        _h1(f"Draft negotiation letter — {contract_name}", plain),
        "",
        f"To: {_RECIPIENT_PLACEHOLDER} · From: {_SENDER_PLACEHOLDER}",
        "",
        f"Our review of the draft {contract_name} against the `{playbook_name}` "
        "playbook found no material issues. Subject to attorney review, we're "
        "prepared to move forward with the current draft.",
        "",
        _disclaimer(plain),
    ]
    return "\n".join(lines).rstrip() + "\n"


def render_letter(
    contract_name: str,
    playbook_name: str,
    findings: list[Finding],
    *,
    recipient: str | None = None,
    sender: str | None = None,
    min_severity: str = "medium",
    fmt: str = "md",
) -> str:
    """Render a draft negotiation letter for ``findings``.

    ``fmt`` is ``"md"`` (default) or ``"txt"`` — plain text for pasting
    straight into an email. Findings below ``min_severity`` are omitted.
    """
    plain = fmt == "txt"
    who_to = recipient or _RECIPIENT_PLACEHOLDER
    who_from = sender or _SENDER_PLACEHOLDER
    selected = select_findings(findings, min_severity)

    if not selected:
        return _render_no_findings(contract_name, playbook_name, plain)

    n = len(selected)
    lines = [
        _h1(f"Draft negotiation letter — {contract_name}", plain),
        "",
        f"To: {who_to} · From: {who_from}",
        "",
        f"Dear {who_to},",
        "",
        f"We've reviewed the draft {contract_name} (checked against the "
        f"`{playbook_name}` playbook) and would like to request the following "
        f"{n} change{'s' if n != 1 else ''} before signing. Each one quotes "
        "the current language, explains our concern, and proposes "
        "specific replacement language where we have it.",
        "",
    ]
    for i, f in enumerate(selected, 1):
        label = _SEVERITY_LABEL.get(f.severity, f.severity.upper())
        lines += [
            _h2(f"{i}. {f.title} — {label}", plain),
            "",
            _bold("The contract says:", plain),
            "",
            _quote(f.excerpt, plain),
            "",
            _bold("Our concern:", plain),
            "",
            f.why.strip(),
            "",
            _bold("Our requested change:", plain),
            "",
            f.suggestion.strip(),
        ]
        if f.fallback.strip():
            lines += [
                "",
                _bold("Proposed language:", plain),
                "",
                _quote(f.fallback, plain),
            ]
        lines.append("")
    lines += [
        "---" if not plain else ("-" * 40),
        "",
        "We'd appreciate a revised draft reflecting these changes, and we're "
        "happy to discuss any of them.",
        "",
        f"Best regards,\n{who_from}",
        "",
        _disclaimer(plain),
    ]
    return "\n".join(lines).rstrip() + "\n"


def _letter_date() -> str:
    today = date.today()
    return f"{today:%B} {today.day}, {today:%Y}"


def _labeled_para(doc: Document, label: str, text: str) -> None:
    """A paragraph whose label is bold, e.g. ``The contract says:`` + text."""
    p = doc.add_paragraph()
    p.add_run(label).bold = True
    if text:
        p.add_run(" " + text.strip())


def _quote_para(doc: Document, text: str) -> None:
    """An indented, italic block quote paragraph for excerpts / proposed language."""
    for line in text.strip().splitlines() or [""]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.add_run(line or " ").italic = True


def _letter_head(
    doc: Document,
    contract_name: str,
    recipient: str | None,
    sender: str | None,
    title: str = "Draft negotiation letter",
) -> tuple[str, str]:
    """Render the title, date, and addressee block. Returns (who_to, who_from)."""
    who_to = recipient or _RECIPIENT_PLACEHOLDER
    who_from = sender or _SENDER_PLACEHOLDER
    doc.add_heading(f"{title} — {contract_name}", level=1)
    doc.add_paragraph(_letter_date())
    doc.add_paragraph(f"To: {who_to}")
    doc.add_paragraph(f"From: {who_from}")
    doc.add_paragraph(f"Dear {who_to},")
    return who_to, who_from


def _letter_disclaimer(doc: Document) -> None:
    doc.add_paragraph().add_run(DISCLAIMER).italic = True


def render_letter_docx(
    contract_name: str,
    playbook_name: str,
    findings: list[Finding],
    *,
    recipient: str | None = None,
    sender: str | None = None,
    min_severity: str = "medium",
) -> bytes:
    """Render the draft negotiation letter as a Word ``.docx``.

    Same content as :func:`render_letter`, formatted as a sendable business
    letter: title, date, addressee block, salutation, numbered change
    requests (severity-labeled, worst first, contract excerpts and proposed
    replacement language as indented quotes), signature, and the
    not-legal-advice disclaimer. Returns the file bytes; a clean review
    produces the short "no material issues" letter.
    """
    doc = Document()
    doc.core_properties.title = f"Draft negotiation letter — {contract_name}"
    doc.core_properties.author = "redline-buddy"

    who_to, who_from = _letter_head(doc, contract_name, recipient, sender)
    selected = select_findings(findings, min_severity)

    if not selected:
        doc.add_paragraph(
            f"Our review of the draft {contract_name} against the "
            f"`{playbook_name}` playbook found no material issues. Subject "
            "to attorney review, we're prepared to move forward with the "
            "current draft."
        )
        _letter_disclaimer(doc)
    else:
        n = len(selected)
        doc.add_paragraph(
            f"We've reviewed the draft {contract_name} (checked against the "
            f"`{playbook_name}` playbook) and would like to request the "
            f"following {n} change{'s' if n != 1 else ''} before signing. "
            "Each one quotes the current language, explains our concern, "
            "and proposes specific replacement language where we have it."
        )
        for i, f in enumerate(selected, 1):
            label = _SEVERITY_LABEL.get(f.severity, f.severity.upper())
            doc.add_heading(f"{i}. {f.title} — {label}", level=2)
            _labeled_para(doc, "The contract says:", "")
            _quote_para(doc, f.excerpt)
            _labeled_para(doc, "Our concern:", f.why)
            _labeled_para(doc, "Our requested change:", f.suggestion)
            if f.fallback.strip():
                _labeled_para(doc, "Proposed language:", "")
                _quote_para(doc, f.fallback)
        doc.add_paragraph(
            "We'd appreciate a revised draft reflecting these changes, and "
            "we're happy to discuss any of them."
        )
        doc.add_paragraph("Best regards,")
        doc.add_paragraph(who_from)
        _letter_disclaimer(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_FOLLOWUP_TAG = {
    "new": "NEW IN THIS DRAFT",
    "redrafted": "STILL FLAGGED AFTER REDRAFTING",
    "outstanding": "NOT ADDRESSED",
}


def classify_followup(
    cmp: Comparison, min_severity: str = "medium"
) -> tuple[list[Finding], list[tuple[Finding, str]]]:
    """Split a round-2 comparison into concessions and remaining issues.

    Returns ``(concessions, remaining)`` where concessions are the resolved
    findings (what the counterparty fixed) and remaining is a list of
    ``(finding, tag)`` for every still-open finding in the new draft:
    ``"new"`` (gained in this draft), ``"redrafted"`` (reworded but still
    flagged), or ``"outstanding"`` (flagged in both drafts, unchanged).
    Both lists are severity-filtered and ordered worst-first.
    """
    concessions = select_findings(cmp.resolved, min_severity)
    gained_ids = {f.rule_id for f in cmp.gained}
    redrafted_ids = {r.after.rule_id for r in cmp.reworded}
    remaining: list[tuple[Finding, str]] = []
    for f in select_findings(cmp.findings_new, min_severity):
        if f.rule_id in gained_ids:
            tag = "new"
        elif f.rule_id in redrafted_ids:
            tag = "redrafted"
        else:
            tag = "outstanding"
        remaining.append((f, tag))
    return concessions, remaining


def _followup_opening(cmp: Comparison) -> str:
    return (
        f"Thanks for the revised draft. We compared {cmp.new_name} against "
        f"{cmp.old_name} (checked against the `{cmp.playbook_name}` playbook)."
    )


def _render_change_request(
    lines: list[str], i: int, f: Finding, tag: str, plain: bool
) -> None:
    label = _SEVERITY_LABEL.get(f.severity, f.severity.upper())
    lines += [
        _h2(f"{i}. {f.title} — {label} · {_FOLLOWUP_TAG[tag]}", plain),
        "",
        _bold("The contract says:", plain),
        "",
        _quote(f.excerpt, plain),
        "",
        _bold("Our concern:", plain),
        "",
        f.why.strip(),
        "",
        _bold("Our requested change:", plain),
        "",
        f.suggestion.strip(),
    ]
    if f.fallback.strip():
        lines += [
            "",
            _bold("Proposed language:", plain),
            "",
            _quote(f.fallback, plain),
        ]
    lines.append("")


def render_followup(
    cmp: Comparison,
    *,
    recipient: str | None = None,
    sender: str | None = None,
    min_severity: str = "medium",
    fmt: str = "md",
) -> str:
    """Draft a round-2 follow-up letter from a :class:`Comparison`.

    Thanks the counterparty for the concessions won, then lists every
    still-open change request tagged NEW IN THIS DRAFT, STILL FLAGGED AFTER
    REDRAFTING, or NOT ADDRESSED. When nothing remains, produces the short
    ready-to-move-forward letter. ``fmt`` is ``"md"`` (default) or ``"txt"``.
    """
    plain = fmt == "txt"
    who_to = recipient or _RECIPIENT_PLACEHOLDER
    who_from = sender or _SENDER_PLACEHOLDER
    concessions, remaining = classify_followup(cmp, min_severity)

    lines = [
        _h1(f"Follow-up letter — {cmp.new_name}", plain),
        "",
        f"To: {who_to} · From: {who_from}",
        "",
        f"Dear {who_to},",
        "",
        _followup_opening(cmp),
        "",
    ]
    if not remaining:
        lines += [
            "Every issue from our earlier review is resolved in this draft. "
            "Subject to attorney review, we're ready to move forward.",
            "",
            f"Best regards,\n{who_from}",
            "",
            _disclaimer(plain),
        ]
        return "\n".join(lines).rstrip() + "\n"

    lines += [
        _h2("What's fixed — thank you", plain),
        "",
    ]
    if concessions:
        lines += [
            "The following issues from our earlier review are resolved in "
            "this draft:"
        ]
        for f in concessions:
            label = _SEVERITY_LABEL.get(f.severity, f.severity.upper())
            lines.append(f"- {f.title} ({label})")
        lines.append("")
    else:
        lines += [
            "This draft doesn't resolve any of the issues from our earlier "
            "review.",
            "",
        ]

    n = len(remaining)
    lines += [
        _h2(f"Still open — {n} change request{'s' if n != 1 else ''}", plain),
        "",
    ]
    for i, (f, tag) in enumerate(remaining, 1):
        _render_change_request(lines, i, f, tag, plain)
    lines += [
        "---" if not plain else ("-" * 40),
        "",
        "We'd appreciate a revised draft addressing the items above, and "
        "we're happy to discuss any of them.",
        "",
        f"Best regards,\n{who_from}",
        "",
        _disclaimer(plain),
    ]
    return "\n".join(lines).rstrip() + "\n"


def render_followup_docx(
    cmp: Comparison,
    *,
    recipient: str | None = None,
    sender: str | None = None,
    min_severity: str = "medium",
) -> bytes:
    """Render the round-2 follow-up letter as a Word ``.docx``.

    Same content as :func:`render_followup`: date, addressee block,
    concessions-won list, then the numbered still-open change requests with
    their NEW / REDRAFTED / NOT ADDRESSED tags. Returns the file bytes.
    """
    doc = Document()
    doc.core_properties.title = f"Follow-up letter — {cmp.new_name}"
    doc.core_properties.author = "redline-buddy"

    who_to, who_from = _letter_head(
        doc, cmp.new_name, recipient, sender, title="Follow-up letter"
    )
    concessions, remaining = classify_followup(cmp, min_severity)

    doc.add_paragraph(_followup_opening(cmp).replace("`", ""))

    if not remaining:
        doc.add_paragraph(
            "Every issue from our earlier review is resolved in this draft. "
            "Subject to attorney review, we're ready to move forward."
        )
        doc.add_paragraph("Best regards,")
        doc.add_paragraph(who_from)
        _letter_disclaimer(doc)
    else:
        doc.add_heading("What's fixed — thank you", level=2)
        if concessions:
            doc.add_paragraph(
                "The following issues from our earlier review are resolved "
                "in this draft:"
            )
            for f in concessions:
                label = _SEVERITY_LABEL.get(f.severity, f.severity.upper())
                doc.add_paragraph(f"{f.title} ({label})", style="List Bullet")
        else:
            doc.add_paragraph(
                "This draft doesn't resolve any of the issues from our "
                "earlier review."
            )
        n = len(remaining)
        doc.add_heading(
            f"Still open — {n} change request{'s' if n != 1 else ''}", level=2
        )
        for i, (f, tag) in enumerate(remaining, 1):
            label = _SEVERITY_LABEL.get(f.severity, f.severity.upper())
            doc.add_heading(f"{i}. {f.title} — {label} · {_FOLLOWUP_TAG[tag]}", level=3)
            _labeled_para(doc, "The contract says:", "")
            _quote_para(doc, f.excerpt)
            _labeled_para(doc, "Our concern:", f.why)
            _labeled_para(doc, "Our requested change:", f.suggestion)
            if f.fallback.strip():
                _labeled_para(doc, "Proposed language:", "")
                _quote_para(doc, f.fallback)
        doc.add_paragraph(
            "We'd appreciate a revised draft addressing the items above, "
            "and we're happy to discuss any of them."
        )
        doc.add_paragraph("Best regards,")
        doc.add_paragraph(who_from)
        _letter_disclaimer(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
